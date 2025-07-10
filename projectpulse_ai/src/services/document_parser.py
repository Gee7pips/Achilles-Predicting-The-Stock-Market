"""Document parsing service.

Very lightweight heuristic extraction for MVP. In production you might switch to Textract, Azure Form Recognizer, etc.
"""
from pathlib import Path
from typing import Tuple, Dict
import re

import PyPDF2
import docx
import openpyxl
import pytesseract

FIELD_REGEXES = {
    "project_name": re.compile(r"project name[:\-]?\s*(.+)", re.IGNORECASE),
    "contractor": re.compile(r"contractor[:\-]?\s*(.+)", re.IGNORECASE),
    "budget": re.compile(r"budget[:\-]?\s*([$\d,\.]+)", re.IGNORECASE),
    "milestones": re.compile(r"milestone[s]?:?\s*(.+)", re.IGNORECASE),
}


def _extract_text_pdf(path: Path) -> str:
    reader = PyPDF2.PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        # Attempt OCR if PDF is scanned images
        text = _ocr_pdf(path)
    return text


def _ocr_pdf(path: Path) -> str:
    try:
        import pdf2image  # lazy import, optional
    except ImportError:
        return ""  # OCR not available
    images = pdf2image.convert_from_path(str(path))
    return "\n".join(pytesseract.image_to_string(img) for img in images)


def _extract_text_docx(path: Path) -> str:
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_text_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    texts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            texts.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
    return "\n".join(texts)


def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_text_pdf(path)
    if ext in {".docx", ".doc"}:
        return _extract_text_docx(path)
    if ext in {".xlsx", ".xls"}:
        return _extract_text_xlsx(path)
    # Fallback to textract for other types
    try:
        import textract  # optional heavy dependency
        return textract.process(str(path)).decode()
    except Exception:
        return ""


def _extract_fields(text: str) -> Dict[str, str]:
    fields: Dict[str, str] = {}
    for key, pattern in FIELD_REGEXES.items():
        match = pattern.search(text)
        if match:
            fields[key] = match.group(1).strip()
    return fields


def parse_document(path: Path) -> Tuple[str, Dict[str, str]]:
    """Return raw text and extracted fields."""
    text = _extract_text(path)
    fields = _extract_fields(text)
    return text, fields


def classify_document(text: str) -> str:
    text_lower = text.lower()
    if "invoice" in text_lower:
        return "invoice"
    if "contract" in text_lower:
        return "contract"
    if "progress report" in text_lower or "report" in text_lower:
        return "report"
    return "unknown"